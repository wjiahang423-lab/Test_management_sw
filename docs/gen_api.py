# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT_DIR = "/home/jyzn/Test_management_sw/docs"

def sf(run):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)

def ah(doc, text, lvl=1):
    h = doc.add_heading(text, level=lvl)
    for r in h.runs:
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.color.rgb = RGBColor(0x1F, 0x5A, 0xA0)
    return h

def ap(doc, text, bp=None):
    p = doc.add_paragraph()
    if bp:
        r = p.add_run(bp); r.font.bold=True; sf(r); p.add_run(text)
    else:
        r = p.add_run(text); sf(r)
    return p

def at(doc, headers, rows, widths=None, hc="1F5AA0"):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = h
        tc = c._tc; tp = tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hc)
        tp.append(sh)
        for pp in c.paragraphs:
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in pp.runs:
                r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                sf(r); r.font.size = Pt(9)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            if ci < len(t.rows[ri+1].cells):
                c = t.rows[ri+1].cells[ci]; c.text = str(val)
                if ri % 2 == 1:
                    tc = c._tc; tp = tc.get_or_add_tcPr()
                    sh = OxmlElement("w:shd")
                    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), "F0F4FA")
                    tp.append(sh)
                for pp in c.paragraphs:
                    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in pp.runs:
                        sf(r); r.font.size = Pt(9)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)
    return t

def ac(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(text); r.font.name = "Consolas"; r.font.size = Pt(8.5)
    return p

doc = Document()
s = doc.styles["Normal"]; s.font.name = "Microsoft YaHei"; s.font.size = Pt(10.5)
s.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

for _ in range(5): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("测试用例管理与执行系统"); r.font.size = Pt(28); r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x5A, 0xA0); sf(r)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("软件接口文档"); r.font.size = Pt(36); r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x5A, 0xA0); sf(r)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("\n\n版本：V1.0\n日期：2026-09-11"); r.font.size = Pt(14); sf(r)
doc.add_page_break()

ah(doc, "1. 接口概述", 1)
ap(doc, "本文档描述测试用例管理与执行系统提供的所有接口，包括测试脚本内置API、计划文件数据格式、数据上报接口。")
at(doc, ["接口类别", "所属模块", "用途"], [
    ["测试脚本API", "context.py / api.py", "供测试脚本调用的内置函数（20+个）"],
    ["计划文件格式", "plan_model.py", ".plan JSON文件的完整结构定义"],
    ["JSON数据上报", "report.py::send_json_cases", "逐用例测试结果POST到服务器"],
    ["MES报告上报", "report.py::send_remote_report", "测试汇总结果POST到MES系统"],
    ["远程文件上传", "report.py::upload_reports", "HTML报告+日志文件multipart上传"],
], widths=[3, 4, 7])

ah(doc, "2. 测试脚本内置 API", 1)
ap(doc, "测试脚本通过 import test_api 或直接使用全局函数名访问以下API。所有函数均线程安全。")

ah(doc, "2.1 SN与结果相关", 2)
at(doc, ["函数签名", "返回值", "说明"], [
    ["set_sn_to_Panel(sn: str) -> bool", "bool", "设置面板显示的SN码"],
    ["get_sn() -> str", "str", "获取当前SN码"],
    ["get_test_result() -> bool|None", "bool|None", "获取当前用例判定结果"],
    ["get_last_test_result() -> bool|None", "bool|None", "获取上一个用例判定结果"],
    ["get_last_test_detail() -> str", "str", "获取上一个用例详情文本"],
], widths=[5, 3, 6])

ah(doc, "2.2 日志相关", 2)
at(doc, ["函数签名", "返回值", "说明"], [
    ["log(msg: str) -> bool", "bool", "输出日志到运行日志和GUI日志区"],
], widths=[5, 3, 6])
ap(doc, "log() 输出格式：[脚本] <msg>，同时写入运行日志文件和GUI日志区。")

ah(doc, "2.3 显示信息与变量", 2)
at(doc, ["函数签名", "返回值", "说明"], [
    ["set_display_info(key: str, value: any) -> bool", "bool", "设置UI显示信息"],
    ["get_display_info(key: str|None) -> dict|any", "dict|any", "获取UI显示信息"],
    ["get_variable(name: str) -> any", "any", "获取全局变量值"],
    ["set_variable(name: str, value: any) -> bool", "bool", "设置全局变量值"],
], widths=[5, 3, 6])

ah(doc, "2.4 机器人状态相关", 2)
at(doc, ["函数签名", "返回值", "说明"], [
    ["set_robot_status(**kwargs) -> bool", "bool", "批量更新机器人状态"],
    ["set_robot_temperature(val: float) -> bool", "bool", "设置温度"],
    ["set_robot_current(val: float) -> bool", "bool", "设置电流"],
    ["set_robot_voltage(val: float) -> bool", "bool", "设置电压"],
    ["set_robot_battery(val: int) -> bool", "bool", "设置电量百分比"],
    ["get_robot_status(key: str|None) -> any", "any", "获取机器人状态"],
], widths=[4.5, 2.5, 5])

ah(doc, "2.5 测量结果相关（Loop专用）", 2)
at(doc, ["函数签名", "返回值", "说明"], [
    ["set_measure_result(value=None, expected=None, unit=None, upper=None, lower=None, message=None, passed=None) -> bool", "bool", "写入测量结果字段"],
    ["set_measure_value(value, unit=None) -> bool", "bool", "快捷设置值和单位"],
    ["set_measure_expected(expected) -> bool", "bool", "快捷设置期望值"],
    ["set_measure_range(lower, upper) -> bool", "bool", "快捷设置阈值范围"],
    ["set_measure_field(key: str, value: any) -> bool", "bool", "写入任意自定义字段"],
    ["get_measure_result() -> dict", "dict", "获取当前测量结果字典"],
    ["reset_measure_result() -> bool", "bool", "清空测量结果"],
], widths=[5.5, 2, 6.5])

ah(doc, "2.6 上报与工位相关", 2)
at(doc, ["函数签名", "返回值", "说明"], [
    ["set_upload_key(key: str) -> bool", "bool", "设置逐用例JSON上报密钥"],
    ["get_upload_key() -> str", "str", "获取上报密钥"],
    ["get_station_id() -> str", "str", "获取工位ID"],
    ["get_station_name() -> str", "str", "获取工位名称"],
], widths=[4.5, 2.5, 5])

ah(doc, "3. 测试计划文件格式 (.plan)", 1)
ap(doc, "测试计划以JSON格式存储，扩展名为 .plan。")

ah(doc, "3.1 顶层结构", 2)
ac(doc, '{\n  "name": "计划名称",\n  "description": "计划描述",\n  "variables": { "name": { "type": "str", "value": "", "description": "" } },\n  "settings": {\n    "storage_mode": "local",\n    "batch": "",\n    "continuous": false,\n    "json_upload_enabled": false,\n    "json_upload_url": "",\n    "json_upload_key": "",\n    "server_username": "root",\n    "server_password": "root"\n  },\n  "sequences": [\n    {\n      "name": "序列名称",\n      "cases": [\n        {\n          "id": "随机UUID",\n          "name": "用例名称",\n          "type": "measurement",\n          "timeout_ms": 8000,\n          "retry": 0,\n          "fail_policy": "continue",\n          "skip": false,\n          "config": { ... }\n        }\n      ]\n    }\n  ]\n}')

ah(doc, "3.2 Measurement用例 config", 2)
ac(doc, '{\n  "script": "scripts/pi_mock/pi_mock.py",\n  "function": "measure_pi",\n  "params": [\n    {"name": "path", "source": "value", "type": "str", "value": "/api/voltage"}\n  ],\n  "returns": [\n    {"item": "value", "judge": "范围内", "threshold": "10~14", "bind_var": ""},\n    {"item": "code", "judge": "等于", "threshold": "0", "bind_var": ""}\n  ]\n}')

ah(doc, "3.3 Loop用例 config", 2)
ac(doc, '{\n  "script": "scripts/auto_api.py",\n  "function": "auto_loop",\n  "parser": "parse_result",\n  "yaml_path": "scripts/demo_loop_data.yaml",\n  "session_key": "sessions",\n  "overrides": [{"item": "voltage", "type": "float", "value": "12.5"}]\n}')

ah(doc, "3.4 其他用例 config", 2)
ap(doc, "Pop: {title, content, btn_true, btn_false, result_var}", bp="")
ap(doc, "Delay: {delay_ms, description}", bp="")
ap(doc, "Action: {script, function}", bp="")

ah(doc, "4. 数据上报接口", 1)

ah(doc, "4.1 逐用例JSON上报", 2)
ap(doc, "POST到 json_upload_url，Body格式：")
ac(doc, '{\n  "key": "上报密钥",\n  "sn": "产品SN",\n  "batch": "批次号",\n  "station_id": "工位ID",\n  "test_time": "2026-09-10 08:30:00.123",\n  "records": [\n    {\n      "test_time": "...",\n      "case_name": "电压测量",\n      "type": "measurement",\n      "status": "PASS",\n      "value": 12.3,\n      "expected": "10~14",\n      "threshold_upper": 14.0,\n      "threshold_lower": 10.0\n    }\n  ]\n}')
ap(doc, "认证：Basic Auth (server_username/server_password)")

ah(doc, "4.2 MES远程报告上报", 2)
ap(doc, "POST到 mes_server/mes_interface，Body包含 plan/sn/overall/total/passed/failed/skipped/results")

ah(doc, "4.3 远程文件上传", 2)
ap(doc, "multipart/form-data上传到 remote_storage_url，字段：report/log/plan/sn")

ah(doc, "5. 日志接口", 1)
ap(doc, "系统日志：data/logs/system_YYYYMMDD.log")
ap(doc, "运行日志：data/logs/<计划名>_<时间戳>.log")
at(doc, ["级别", "前缀", "来源"], [
    ["INFO", "[INFO]", "正常运行信息"],
    ["WARN", "[WARN]", "警告信息（含Qt解析警告）"],
    ["ERROR", "[ERROR]", "错误信息（含traceback）"],
    ["DEBUG", "[DEBUG]", "调试信息"],
    ["RUN", "[RUN]", "运行日志特有前缀"],
], widths=[3, 3, 8])
ap(doc, "日志行格式：[YYYY-MM-DD HH:MM:SS] [LEVEL] message")

ah(doc, "6. 完整计划文件示例", 1)
ac(doc, '''{
  "name": "设备与执行页面接口巡检（连续）",
  "description": "以树莓派mock服务为服务器做接口巡检",
  "variables": {},
  "settings": {
    "storage_mode": "local",
    "continuous": true,
    "json_upload_enabled": false
  },
  "sequences": [{
    "name": "接口巡检序列",
    "cases": [
      {
        "id": "abc123",
        "name": "设备状态Robot",
        "type": "measurement",
        "timeout_ms": 8000,
        "retry": 0,
        "fail_policy": "continue",
        "config": {
          "script": "scripts/pi_mock/pi_mock.py",
          "function": "get_robot_status",
          "params": [{"name":"base","source":"value","type":"str","value":"http://10.5.35.49:5000"}],
          "returns": [
            {"item":"temperature","judge":"范围内","threshold":"10~40","bind_var":""},
            {"item":"current","judge":"范围内","threshold":"0~3","bind_var":""},
            {"item":"voltage","judge":"范围内","threshold":"10~14","bind_var":""},
            {"item":"battery","judge":"大于","threshold":"50","bind_var":""},
            {"item":"code","judge":"等于","threshold":"0","bind_var":""}
          ]
        }
      },
      {
        "id": "def456",
        "name": "延时500ms",
        "type": "delay",
        "timeout_ms": 6000,
        "config": {"delay_ms": 500}
      }
    ]
  }]
}''')

sec = doc.sections[0]
footer = sec.footer; fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("测试用例管理与执行系统  软件接口文档  V1.0  2026-09-11")
r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x99,0x99,0x99); sf(r)

doc.save(os.path.join(OUTPUT_DIR, "软件接口文档.docx"))
print("OK: 软件接口文档.docx")
