# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT_DIR = "/home/jyzn/Test_management_sw/docs"

def sf(run):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)

def ah(doc, text, lvl=1):
    h = doc.add_heading(text, level=lvl)
    for r in h.runs:
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.color.rgb = RGBColor(0x1F, 0x5A, 0xA0)
    return h

def ap(doc, text, bp=None):
    p = doc.add_paragraph()
    if bp:
        r = p.add_run(bp); r.font.bold=True; sf(r); p.add_run(text)
    else:
        r = p.add_run(text); sf(r)
    return p

def at(doc, headers, rows, widths=None, hc="1F5AA0"):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = h
        tc = c._tc; tp = tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hc)
        tp.append(sh)
        for pp in c.paragraphs:
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in pp.runs:
                r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                sf(r); r.font.size = Pt(9)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            if ci < len(t.rows[ri+1].cells):
                c = t.rows[ri+1].cells[ci]; c.text = str(val)
                if ri % 2 == 1:
                    tc = c._tc; tp = tc.get_or_add_tcPr()
                    sh = OxmlElement("w:shd")
                    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), "F0F4FA")
                    tp.append(sh)
                for pp in c.paragraphs:
                    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in pp.runs:
                        sf(r); r.font.size = Pt(9)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)
    return t

def ac(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(text); r.font.name = "Consolas"; r.font.size = Pt(8.5)
    return p

doc = Document()
s = doc.styles["Normal"]; s.font.name = "Microsoft YaHei"; s.font.size = Pt(10.5)
s.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

for _ in range(5): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("测试用例管理与执行系统"); r.font.size = Pt(28); r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x5A, 0xA0); sf(r)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("软件说明文档"); r.font.size = Pt(36); r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x5A, 0xA0); sf(r)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("\n\n版本：V1.0\n日期：2026-09-11"); r.font.size = Pt(14); sf(r)
doc.add_page_break()

ah(doc, "1. 软件简介", 1)
ap(doc, "本软件是一套面向产线/实验室的测试用例管理与执行平台，专为机器人EOL测试场景设计。支持五种用例类型、连续循环测试模式、远程数据上报和可视化报告生成，可在产线平板上以Kiosk全屏模式7x24小时稳定运行。")
at(doc, ["功能特性", "说明"], [
    ["多级计划管理", "测试计划->测试序列->测试用例，三层结构"],
    ["五种用例类型", "Action/Delay/Pop/Measurement/Loop"],
    ["权限控制", "管理员/操作员两级角色，一键分权登录"],
    ["连续测试模式", "continuous=true时自动循环执行"],
    ["生产统计看板", "总数量/良品数/不良品数/良品率/不良品率"],
    ["HTML报告", "每轮测试自动生成HTML报告"],
    ["数据上报", "支持逐用例JSON上报和MES远程报告上传"],
    ["脚本API", "测试脚本可直接调用20+内置API函数"],
    ["Kiosk模式", "无边框全屏，状态栏显示WiFi/电量"],
], widths=[4, 10])

ah(doc, "2. 运行环境", 1)
at(doc, ["项目", "要求"], [
    ["操作系统", "Windows 10+ / Ubuntu 20.04+ / Raspberry Pi OS"],
    ["Python版本", "3.8+（开发环境3.10）"],
    ["内存", ">=4GB（推荐8GB）"],
    ["磁盘", ">=2GB可用空间"],
    ["显示", ">=1280x800分辨率（触屏推荐）"],
], widths=[4, 10])
ap(doc, "依赖包：PyQt5>=5.15, PyYAML>=5.1, requests>=2.20, paramiko>=3.0, zlgcan")

ah(doc, "3. 安装与启动", 1)
ah(doc, "3.1 源码运行", 2)
ac(doc, "cd Test_management_sw\npip install -r requirements.txt\npython3 main.py")
ah(doc, "3.2 打包部署", 2)
ac(doc, "python3 build_app.py\n# 输出: dist/EOL/")
ah(doc, "3.3 远程部署", 2)
ap(doc, "到树莓派：", bp="")
ac(doc, "./deploy/push_to_pi.sh 10.5.35.49 pi raspberry on")
ap(doc, "到工业平板：", bp="")
ac(doc, "scp -r dist/EOL/ user@10.5.33.219:/home/user/eol/\nssh user@10.5.33.219 'systemctl restart eol'")

ah(doc, "4. 功能说明", 1)
ah(doc, "4.1 执行页面", 2)
at(doc, ["区域", "功能"], [
    ["顶部标题栏", "Logo + 工位名称 + 时钟"],
    ["左侧步骤树", "测试序列和用例列表，实时更新状态"],
    ["右侧信息区", "SN显示、批次、测试节拍、进度条"],
    ["生产统计区", "总数量/良品数/不良品数/良品率/不良品率"],
    ["执行日志", "可展开/折叠，最多5000行"],
    ["开始按钮", "开始 / 暂停 / 停止"],
], widths=[3.5, 10.5])
ah(doc, "4.2 管理页面（仅管理员）", 2)
at(doc, ["子页面", "功能"], [
    ["测试计划", "新建/打开/保存/删除测试计划"],
    ["全局变量", "管理系统级变量"],
    ["系统设置", "字体/尺寸/速度/密码/工位信息"],
    ["用户管理", "添加/修改/删除用户"],
], widths=[3, 10])

ah(doc, "5. 测试计划编辑指南", 1)
ah(doc, "5.1 用例配置说明", 2)
at(doc, ["用例类型", "配置项", "说明"], [
    ["Action", "脚本、函数", "无输入参数，仅执行"],
    ["Delay", "延时(ms)、说明", "固定等待指定时间"],
    ["Pop", "标题、内容、按钮文本", "弹出确认框"],
    ["Measurement", "脚本、函数、参数、返回判定", "有参数，返回值与阈值比较"],
    ["Loop", "脚本、函数、YAML数据源", "批量执行，每个session独立判定"],
], widths=[3, 4, 7])
ah(doc, "5.2 Measurement判定配置", 2)
at(doc, ["返回项(item)", "判定(judge)", "阈值(threshold)", "示例"], [
    ["value", "范围内", "10~14", "电压10-14V通过"],
    ["value", "等于", "0", "code=0通过"],
    ["temperature", "范围内", "10~40", "温度10-40C通过"],
    ["battery", "大于", "50", "电量>50%通过"],
], widths=[3, 3, 3, 5])

ah(doc, "6. 快捷键说明", 1)
at(doc, ["快捷键", "功能"], [
    ["F11", "切换全屏/窗口模式"],
    ["Ctrl+M", "管理员切换到管理页面"],
    ["Ctrl+Shift+L", "重新登录"],
    ["Ctrl+Q", "退出程序"],
], widths=[4, 10])

ah(doc, "7. 故障排查", 1)
at(doc, ["症状", "可能原因", "排查方法"], [
    ["软件闪退", "display未设置/Qt插件缺失", "检查DISPLAY环境变量，查看system日志"],
    ["脚本加载失败", "路径错误/缺少依赖", "检查.plan中脚本路径"],
    ["用例一直FAIL", "mock服务不可达/阈值错误", "ping mock服务器，检查threshold配置"],
    ["统计不归零", "密码错误", "确认clear_password设置值"],
    ["字体溢出边框", "字体过大+UI padding不足", "调小exec_small_font或font_size"],
], widths=[3.5, 4, 6.5])

sec = doc.sections[0]
footer = sec.footer; fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("测试用例管理与执行系统  软件说明文档  V1.0  2026-09-11")
r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x99,0x99,0x99); sf(r)

doc.save(os.path.join(OUTPUT_DIR, "软件说明.docx"))
print("OK: 软件说明.docx")
