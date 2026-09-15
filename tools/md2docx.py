#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown 转 DOCX 转换器（基于 python-docx）"""
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


def set_cell_shading(cell, color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def add_code_block(doc, code_text):
    """添加代码块（灰底等宽字体）"""
    for line in code_text.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(14)
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x24, 0x29, 0x2e)
        # 灰底
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F6F8FA"/>')
        run._element.get_or_add_rPr().append(shading)


def add_table(doc, headers, rows):
    """添加表格"""
    cols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 表头
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h.strip())
        run.bold = True
        run.font.size = Pt(9)
        set_cell_shading(cell, "E8ECF0")
    # 数据行
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            val = row[c_idx] if c_idx < len(row) else ""
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val.strip())
            run.font.size = Pt(9)
    doc.add_paragraph()  # 表后空行


def parse_inline(text):
    """解析行内格式（粗体、代码）"""
    parts = []
    pattern = r'(\*\*(.+?)\*\*|`([^`]+)`|([^*`]+))'
    for m in re.finditer(pattern, text):
        if m.group(2):  # bold
            parts.append(("bold", m.group(2)))
        elif m.group(3):  # code
            parts.append(("code", m.group(3)))
        elif m.group(4):  # plain
            parts.append(("plain", m.group(4)))
    return parts if parts else [("plain", text)]


def add_rich_paragraph(doc, text, style=None, indent_level=0):
    """添加带行内格式的段落"""
    p = doc.add_paragraph(style=style)
    if indent_level > 0:
        p.paragraph_format.left_indent = Cm(indent_level * 0.8)
    parts = parse_inline(text)
    for fmt, content in parts:
        run = p.add_run(content)
        run.font.size = Pt(10.5)
        if fmt == "bold":
            run.bold = True
        elif fmt == "code":
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F6F8FA"/>')
            run._element.get_or_add_rPr().append(shading)
    return p


def md_to_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()
    doc.styles["Normal"].font.name = "Microsoft YaHei"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    # 设置页边距
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    i = 0
    in_code_block = False
    code_lines = []
    in_table = False
    table_headers = []
    table_rows = []

    while i < len(lines):
        line = lines[i].rstrip("\n")

        # 代码块处理
        if line.startswith("```"):
            if in_code_block:
                add_code_block(doc, "\n".join(code_lines))
                code_lines = []
                in_code_block = False
            else:
                # 如果之前在收集表格，先输出
                if in_table and table_headers:
                    add_table(doc, table_headers, table_rows)
                    table_headers = []
                    table_rows = []
                    in_table = False
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # 表格处理
        if "|" in line and line.strip().startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            # 检查是否是分隔行
            if all(re.match(r'^[-:]+$', c) for c in cells):
                i += 1
                continue
            if not in_table:
                table_headers = cells
                in_table = True
            else:
                table_rows.append(cells)
            i += 1
            continue
        else:
            if in_table and table_headers:
                add_table(doc, table_headers, table_rows)
                table_headers = []
                table_rows = []
                in_table = False

        stripped = line.strip()

        # 空行
        if not stripped:
            i += 1
            continue

        # 水平线
        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            # 添加底部边框模拟水平线
            pPr = p._element.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/></w:pBdr>')
            pPr.append(pBdr)
            i += 1
            continue

        # 标题
        heading_match = re.match(r'^(#{1,6})\s+(.+)$', stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        # 无序列表
        list_match = re.match(r'^(\s*)[-*]\s+(.+)$', stripped)
        if list_match:
            indent = len(line) - len(line.lstrip())
            level = indent // 2
            text = list_match.group(2)
            add_rich_paragraph(doc, "  " * level + "  \u2022  " + text)
            i += 1
            continue

        # 有序列表
        ol_match = re.match(r'^(\s*)\d+\.\s+(.+)$', stripped)
        if ol_match:
            text = ol_match.group(2)
            add_rich_paragraph(doc, "    \u2022  " + text)
            i += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            text = stripped.lstrip("> ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            pPr = p._element.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="12" w:space="4" w:color="CCCCCC"/></w:pBdr>')
            pPr.append(pBdr)
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            i += 1
            continue

        # 普通段落
        add_rich_paragraph(doc, stripped)
        i += 1

    # 输出最后残留的表格
    if in_table and table_headers:
        add_table(doc, table_headers, table_rows)

    doc.save(docx_path)
    print(f"已生成: {docx_path}")


if __name__ == "__main__":
    md_file = sys.argv[1] if len(sys.argv) > 1 else "docs/测试脚本使用与接口文档.md"
    docx_file = sys.argv[2] if len(sys.argv) > 2 else md_file.replace(".md", ".docx")
    md_to_docx(md_file, docx_file)
