# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT_DIR = "/home/jyzn/Test_management_sw/docs"

def sf(run):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)

def ah(doc, text, lvl=1):
    h = doc.add_heading(text, level=lvl)
    for r in h.runs:
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.color.rgb = RGBColor(0x1F, 0x5A, 0xA0)
    return h

def ap(doc, text, bp=None):
    p = doc.add_paragraph()
    if bp:
        r = p.add_run(bp); r.font.bold=True; sf(r); p.add_run(text)
    else:
        r = p.add_run(text); sf(r)
    return p

def at(doc, headers, rows, widths=None, hc="1F5AA0"):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = h
        tc = c._tc; tp = tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hc)
        tp.append(sh)
        for pp in c.paragraphs:
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in pp.runs:
                r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                sf(r); r.font.size = Pt(9)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            if ci < len(t.rows[ri+1].cells):
                c = t.rows[ri+1].cells[ci]; c.text = str(val)
                if ri % 2 == 1:
                    tc = c._tc; tp = tc.get_or_add_tcPr()
                    sh = OxmlElement("w:shd")
                    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), "F0F4FA")
                    tp.append(sh)
                for pp in c.paragraphs:
                    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in pp.runs:
                        sf(r); r.font.size = Pt(9)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)
    return t

def ac(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(text); r.font.name = "Consolas"; r.font.size = Pt(8.5)
    return p

doc = Document()
s = doc.styles["Normal"]; s.font.name = "Microsoft YaHei"; s.font.size = Pt(10.5)
s.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

for _ in range(5): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("测试用例管理与执行系统"); r.font.size = Pt(28); r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x5A, 0xA0); sf(r)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("架构设计说明书"); r.font.size = Pt(36); r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x5A, 0xA0); sf(r)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("\n\n版本：V1.0\n日期：2026-09-11"); r.font.size = Pt(14); sf(r)
doc.add_page_break()

ah(doc, "1. 系统概述", 1)
ap(doc, "本系统是一套面向产线/实验室的测试用例管理与执行平台，采用 PyQt5 构建图形界面，通过异步线程模型实现测试用例的非阻塞执行。系统支持五种用例类型，提供完整的计划管理、脚本执行、结果上报和统计展示功能。")
at(doc, ["层级", "技术选型", "说明"], [
    ["UI框架", "PyQt5 (Qt5.15+)", "跨平台桌面GUI，Fusion风格，Kiosk全屏模式"],
    ["执行引擎", "QThread + concurrent.futures", "异步线程池，最多4并发Worker"],
    ["数据模型", "JSON + configparser", "计划/用户/变量/统计均以文件持久化"],
    ["日志系统", "线程安全文件写入", "RLock保护，每日独立文件"],
    ["打包工具", "PyInstaller onedir", "自带PyQt5，无需目标机器安装Python"],
    ["外部依赖", "requests/yaml/paramiko/zlgcan", "网络请求/YAML解析/SSH/CAN总线"],
], widths=[3, 5, 10])

ah(doc, "2. 整体架构", 1)
ap(doc, "系统采用经典的三层架构：表示层（UI）-> 业务逻辑层（Core）-> 数据层（File/Network）。各层之间通过信号槽机制解耦，确保GUI线程与执行线程的隔离。")

ah(doc, "2.1 分层架构图", 2)
chart_lines = [
    "+-------------------------------------------------------------+",
    "|                      表示层 (UI)                             |",
    "|  +----------------+  +----------------+  +------------------+ |",
    "|  |  MainWindow    |  |  ExecutePage   |  |   ManagePage     | |",
    "|  +-------+--------+  +-------+--------+  +--------+---------+ |",
    "|          |                  |                    |            |",
    "|  +-------+------------------+--------------------+-----------+ |",
    "|  |              Qt Signals / Slots 通信机制                   | |",
    "|  +---------------------------------+------------------------+ |",
    "+----------------------------------|----------------------------+",
    "|                      业务逻辑层 (Core)                          |",
    "|  +----------------+  +----------------+  +------------------+  |",
    "|  | EngineWorker   |  |  RuntimeCtx    |  |  PlanModel       |  |",
    "|  +-------+--------+  +-------+--------+  +--------+---------+  |",
    "|  +-------+------------------+--------------------+-----------+ |",
    "|  |  ScriptLoader | ReportBuilder | UserManager               | |",
    "|  +---------------------------------+------------------------+ |",
    "+----------------------------------|----------------------------+",
    "|                      数据层 (Data)                              |",
    "|  +----------+  +----------+  +----------+  +----------------+  |",
    "|  | .plan    |  | JSON     |  | INI      |  | HTML           |  |",
    "|  +----------+  +----------+  +----------+  +----------------+  |",
    "|  +----------+  +----------+  +-----------------------------+  |",
    "|  | HTTP上报 |  | CAN总线  |  | 扫码枪(HID/串口)             |  |",
    "|  +----------+  +----------+  +-----------------------------+  |",
    "+-------------------------------------------------------------+",
]
for line in chart_lines:
    p = doc.add_paragraph()
    r = p.add_run(line)
    r.font.name = "Consolas"; r.font.size = Pt(7.5)

ah(doc, "3. 模块设计", 1)
ah(doc, "3.1 核心模块体系", 2)
at(doc, ["模块文件", "类/函数", "职责"], [
    ["engine.py", "EngineWorker (QThread)", "测试执行引擎核心，管理用例执行、重试、超时、暂停/停止"],
    ["", "_handle_action/delay/pop/measurement/loop", "5种用例类型的具体执行逻辑"],
    ["", "_judge_measurement / _judge_loop_session", "结果判定算法：7种判定模式+自适应传参"],
    ["context.py", "RuntimeContext", "跨线程共享状态容器，所有getter/setter均持RLock"],
    ["", "build_api()", "构建test_api命名空间，供测试脚本调用"],
    ["plan_model.py", "TestCase / TestSequence / TestPlan", "三层数据模型：计划->序列->用例"],
    ["script_loader.py", "load_script_module / load_function", "带mtime缓存的动态脚本加载"],
    ["report.py", "RunLogger / ReportBuilder", "运行日志写入 + HTML报告生成"],
    ["", "send_json_cases / send_remote_report / upload_reports", "三种远程数据上报方式"],
    ["syslog.py", "_write / exception / log_exception_hook", "全局系统日志，RLock保护"],
    ["api.py", "register_api / inject_into_module", "test_api模块注入机制"],
    ["settings_manager.py", "SettingsManager", "全局设置读写"],
    ["user_manager.py", "UserManager", "用户认证管理，支持一键分权"],
    ["variable_manager.py", "VariableManager", "类型化变量存储int/str/float/list/dict"],
    ["paths.py", "resolve_root_path / relpath_from_root", "路径解析与归一化"],
], widths=[3.5, 4.5, 10])

ah(doc, "4. 线程模型", 1)
ap(doc, "系统采用单GUI线程 + 单EngineWorker线程 + ThreadPoolExecutor(4 workers)的三级线程模型。")
at(doc, ["线程", "职责", "同步机制"], [
    ["GUI线程 (主线程)", "所有Qt控件操作、事件处理、信号响应", "QTimer定时刷新"],
    ["EngineWorker (QThread)", "测试用例执行主循环、脚本调用调度", "threading.Event (pause/stop/sn/pop)"],
    ["ThreadPoolExecutor x 4", "并行执行多个被测脚本函数", "concurrent.futures.Future"],
], widths=[4, 7, 7])

ah(doc, "5. 信号体系", 1)
at(doc, ["信号名", "参数", "触发时机", "GUI处理方法"], [
    ["sig_log", "(str)", "每条执行日志", "append_log()"],
    ["sig_case_state", "(case_id, state, detail, elapsed)", "用例状态变更", "_on_case_state()"],
    ["sig_case_sessions", "(case_id, sessions)", "Loop用例session结果", "_on_case_sessions()"],
    ["sig_overall", "(str)", "总体结果 PASS/FAIL", "_on_overall()"],
    ["sig_progress", "(done, total)", "进度更新", "_on_progress()"],
    ["sig_run_finished", "(ok, message)", "测试全部完成", "_on_run_finished()"],
    ["sig_request_pop", "(dict config)", "需要人机交互", "_on_pop_requested()"],
    ["sig_phase", "(str phase)", "运行阶段变更", "_on_phase()"],
    ["sig_round_finished", "(passed, round_no)", "连续模式每轮结束", "_on_round_finished()"],
    ["sig_round_started", "(round_no)", "连续模式每轮开始", "_on_round_started()"],
], widths=[3, 3.5, 3.5, 4])

ah(doc, "6. 关键算法", 1)
ah(doc, "6.1 Measurement判定算法", 2)
at(doc, ["判定模式", "阈值格式", "比较逻辑"], [
    ["等于", "\"5\" 或 \"{a:1}\"(JSON)", "数值或字符串相等"],
    ["不等于", "\"!=5\"", "与等于取反"],
    ["大于/小于", "\">100\" 或 \"<0.5\"", "浮点数比较"],
    ["范围内", "\"10~30\" 或 \"10,30\" 或 \"10，30\"", "lo <= value <= hi"],
    ["包含", "\"hello\"", "字符串包含或list查找"],
    ["长度", "\">3\" 或 \"<=10\" 或 \"3~8\"", "len(value) 比较"],
], widths=[3, 4, 7])

ah(doc, "6.2 Loop参数自适应算法", 2)
at(doc, ["场景", "传参策略", "示例"], [
    ["函数无入参", "func()", "def init() -> dict"],
    ["key完全匹配形参", "func(**item)", "item={v:10}, func(v)"],
    ["部分key匹配", "func(**交集)", "只传匹配的形参"],
    ["函数仅1形参", "func(item)", "def measure(params)"],
    ["函数带**kwargs", "func(**item)", "接受任意键"],
], widths=[4, 5, 5])

ah(doc, "7. 部署架构", 1)
at(doc, ["部署目标", "IP地址", "部署方式", "启动方式"], [
    ["工控机/开发机", "本机", "python3 build_app.py -> dist/EOL/", "双击EOL或python3 main.py"],
    ["树莓派(CIT)", "10.5.35.49", "./deploy/push_to_pi.sh", "systemd自动启动"],
    ["工业平板S101", "10.5.33.219", "scp + systemctl", "eol.service自启"],
], widths=[4, 3, 4, 3])

sec = doc.sections[0]
footer = sec.footer; fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("测试用例管理与执行系统  架构设计说明书  V1.0  2026-09-11")
r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x99,0x99,0x99); sf(r)

doc.save(os.path.join(OUTPUT_DIR, "架构设计.docx"))
print("OK: 架构设计.docx")
